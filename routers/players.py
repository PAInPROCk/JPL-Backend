from fastapi import APIRouter, UploadFile, File, HTTPException, Request, Form
from auth.auth_handler import verify_token, get_token_from_request
from core.database import get_db_connection
import pymysql
import os
import io
import uuid
import zipfile
import csv
import tempfile
import shutil


from typing import List, Optional


router = APIRouter()

UPLOAD_FOLDER_PLAYERS = "uploads/players"

ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}

@router.get("/" \
"players")
def get_players():
    conn = get_db_connection()

    if conn is None:
        return{"error": "Database Connection Failed"}
    
    try:
        cursor  = conn.cursor(pymysql.cursors.DictCursor)

        cursor.execute(
            """SELECT 
    p.id AS player_id,
    p.name,
    p.nickname,
    p.jersey,
    p.category,
    p.type,
    p.image_path,
    p.base_price,
    p.total_runs,
    p.highest_runs,
    p.wickets_taken,
    p.times_out,
    COALESCE(string_agg(DISTINCT t.name, ', ' ORDER BY t.name), '') AS teams_played,
    'player' AS role
FROM players p
LEFT JOIN player_teams pt ON p.id = pt.player_id
LEFT JOIN teams t ON pt.team_id = t.team_id
GROUP BY p.id

UNION ALL

SELECT 
    c.id AS player_id,
    c.name,
    NULL AS nickname,
    NULL AS jersey,
    'Captain' AS category,
    NULL AS type,
    c.image_path,
    NULL AS base_price,
    NULL AS total_runs,
    NULL AS highest_runs,
    NULL AS wickets_taken,
    NULL AS times_out,
    t.name AS teams_played,
    'captain' AS role
FROM captains c
LEFT JOIN teams t ON c.team_id = t.team_id

ORDER BY name;"""
        )

        rows = cursor.fetchall()

        return{
            "success": True,
            "count": len(rows),
            "players": rows
        }
    
    except Exception as e:
        print("Player route error:", e)
        return{"error": str(e)}
    
    finally:
        cursor.close()
        conn.close()

@router.get("/player-with-teams")
def players_with_teams():

    conn = get_db_connection()

    if conn is None:
        return {"error":"Database Connection Failed"}
    
    try:
        cursor = conn.cursor()

        cursor.execute(
            """
                SELECT 
                p.id AS player_id,
                p.name,
                p.nickname,
                p.jersey,
                p.category,
                p.type,
                p.image_path,
                p.base_price,
                 string_agg(t.name, ', ') AS teams_played
            FROM players p
            LEFT JOIN player_teams pt ON p.id = pt.player_id
            LEFT JOIN teams t ON pt.team_id = t.team_id
            GROUP BY p.id
            ORDER BY p.name ASC
            """
        )

        rows = cursor.fetchall()

        return{
            "success": True,
            "count": len(rows),
            "players": rows
        }
    
    except Exception as e:
        print("player-with-teams route error:", e)
        return{"error": str(e)}
    
    finally:
        cursor.close()
        conn.close()

@router.get("/players/{player_id}")
async def get_player(player_id: int, role: str):

    conn = get_db_connection()
    cursor = conn.cursor(pymysql.cursors.DictCursor)

    try:

        if role == "captain":
            cursor.execute(
                "SELECT id, name, team_id, image_path FROM captains WHERE id=%s",
                (player_id,)
            )
            captain = cursor.fetchone()

            if not captain:
                raise HTTPException(status_code=404, detail="Captain not found")

            return {
                "type": "captain",
                "data": captain
            }

        # default → player
        cursor.execute(
            "SELECT * FROM players WHERE id=%s",
            (player_id,)
        )

        player = cursor.fetchone()

        if not player:
            raise HTTPException(status_code=404, detail="Player not found")

        return {
            "type": "player",
            "data": player
        }

    finally:
        cursor.close()
        conn.close()

@router.post("/upload-player-image")
async def upload_player_image(
    request: Request,
    image: UploadFile = File(...)
):

    token = get_token_from_request(request)
    if not token:
        raise HTTPException(status_code=401, detail="Authentication token required")

    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="Invalid or expired token")

    if payload.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin privileges required")

    if not image.filename:
        raise HTTPException(status_code=400, detail="No file provided")

    img_content = await image.read()
    
    # Validate file size (max 5MB) and format
    from core.image_handler import validate_image_bytes
    validate_image_bytes(img_content)
    
    try:
        from PIL import Image
        from core.image_handler import crop_and_resize_to_square, compress_to_webp, upload_image_to_supabase
        
        pil_image = Image.open(io.BytesIO(img_content))
        processed_image = crop_and_resize_to_square(pil_image, 500)
        webp_bytes = compress_to_webp(processed_image)
        
        filename = f"{uuid.uuid4().hex}.webp"
        storage_path = f"players/{filename}"
        public_url = upload_image_to_supabase(webp_bytes, storage_path)
        
        return {
            "image_path": public_url
        }
    except HTTPException:
        raise
    except Exception as e:
        print("❌ Error processing/uploading player image:", e)
        raise HTTPException(status_code=500, detail=f"Image upload failed: {str(e)}")




@router.post("/add-player")
async def add_player(
    request: Request,

    # -------- FORM FIELDS --------
    playerName: Optional[str] = Form(None),
    fatherName: Optional[str] = Form(None),
    surName: Optional[str] = Form(None),
    nickName: Optional[str] = Form(None),
    age: Optional[int] = Form(None),
    category: Optional[str] = Form(None),
    style: Optional[str] = Form(None),
    basePrice: Optional[float] = Form(0),
    totalRuns: Optional[int] = Form(0),
    highestRuns: Optional[int] = Form(0),
    wickets: Optional[int] = Form(0),
    outs: Optional[int] = Form(0),
    jerseyNo: Optional[int] = Form(None),
    mobile: Optional[str] = Form(None),
    emailId: Optional[str] = Form(None),
    gender: Optional[str] = Form(None),

    # -------- MULTI-SELECT TEAMS --------
    teams: Optional[List[int]] = Form([]),

    # -------- FILE --------
    image: Optional[UploadFile] = File(None)
):
    # ================= AUTH =================
    token = get_token_from_request(request)

    if not token:
        raise HTTPException(status_code=401, detail="Unauthorized")

    user = verify_token(token)

    if not user or user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")

    # ================= NAME =================
    full_name = " ".join(
        part for part in [playerName, fatherName, surName] if part
    ).strip()

    if not full_name:
        raise HTTPException(status_code=400, detail="Player full name is required")

    # ================= IMAGE UPLOAD =================
    image_path = None

    if image:
        img_content = await image.read()
        from core.image_handler import validate_image_bytes
        validate_image_bytes(img_content)

        try:
            from PIL import Image
            from core.image_handler import crop_and_resize_to_square, compress_to_webp, upload_image_to_supabase
            
            pil_image = Image.open(io.BytesIO(img_content))
            processed_image = crop_and_resize_to_square(pil_image, 500)
            webp_bytes = compress_to_webp(processed_image)
            
            filename = f"{uuid.uuid4().hex}.webp"
            storage_path = f"players/{filename}"
            image_path = upload_image_to_supabase(webp_bytes, storage_path)
        except HTTPException:
            raise
        except Exception as e:
            print("❌ Error processing/uploading player image in add_player:", e)
            raise HTTPException(status_code=500, detail=f"Image upload failed: {str(e)}")

    # ================= DB =================
    conn = get_db_connection()
    if conn is None:
        raise HTTPException(status_code=500, detail="Database connection failed")

    cursor = conn.cursor(pymysql.cursors.DictCursor)

    try:
        # -------- INSERT PLAYER --------
        cursor.execute("""
            INSERT INTO players 
            (name, nickname, age, category, type, base_price, total_runs, highest_runs, 
             wickets_taken, times_out, image_path, jersey, mobile_no, email_id)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            RETURNING id
        """, (
            full_name,
            nickName,
            age,
            category,
            style,
            basePrice,
            totalRuns,
            highestRuns,
            wickets,
            outs,
            image_path,
            jerseyNo,
            mobile,
            emailId
        ))

        inserted_player = cursor.fetchone()
        player_id = inserted_player["id"]

        # -------- INSERT PLAYER-TEAMS --------
        for team_id in teams:
            cursor.execute(
                "INSERT INTO player_teams (player_id, team_id) VALUES (%s, %s)",
                (player_id, int(team_id))
            )

        conn.commit()

        return {
            "message": "Player added successfully!",
            "player_id": player_id
        }

    except pymysql.IntegrityError as e:
        conn.rollback()
        print("[DB Error] IntegrityError in add_player:", e)
        raise HTTPException(
            status_code=400,
            detail="Player with same name or jersey number exists"
        )

    except Exception as e:
        conn.rollback()
        print("❌ add-player error:", e)
        raise HTTPException(status_code=500, detail=str(e))

    finally:
        cursor.close()
        conn.close()


@router.post("/upload-players")
async def upload_players(request: Request, file: UploadFile = File(...)):

    token = get_token_from_request(request)
    if not token:
        raise HTTPException(401, "Unauthorized")

    payload = verify_token(token)
    if not payload or payload.get("role") != "admin":
        raise HTTPException(403, "Forbidden")

    # ---------- TEMP DIR ----------
    temp_dir = tempfile.mkdtemp()
    zip_path = os.path.join(temp_dir, file.filename)

    # Save ZIP
    with open(zip_path, "wb") as f:
        f.write(await file.read())

    # ---------- EXTRACT ZIP ----------
    try:
        with zipfile.ZipFile(zip_path, 'r') as zip_ref:
            zip_ref.extractall(temp_dir)
    except:
        raise HTTPException(400, "Invalid ZIP file")

    # ---------- FIND EXCEL ----------
    excel_file = None
    for f in os.listdir(temp_dir):
        if f.endswith(".xlsx") or f.endswith(".xls"):
            excel_file = os.path.join(temp_dir, f)

    if not excel_file:
        raise HTTPException(400, "Excel file not found in ZIP")

    # ---------- READ EXCEL ----------
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        sheet = wb.active

        # Extract headers (first row)
        headers = [cell.value for cell in sheet[1]]

        records = []
        for r_idx in range(2, sheet.max_row + 1):
            row_vals = [sheet.cell(row=r_idx, column=c_idx).value for c_idx in range(1, len(headers) + 1)]
            
            # Skip empty rows
            if not any(row_vals):
                continue
                
            record = {}
            for header, val in zip(headers, row_vals):
                if header:
                    record[header] = val
            records.append(record)
    except Exception as exc:
        raise HTTPException(400, f"Error reading Excel file: {exc}")

    # ---------- PROCESS AND UPLOAD IMAGES TO SUPABASE ----------
    images_folder = os.path.join(temp_dir, "images")
    from PIL import Image
    from core.image_handler import crop_and_resize_to_square, compress_to_webp, upload_image_to_supabase

    # ---------- DB INSERT ----------
    conn = get_db_connection()
    if conn is None:
        raise HTTPException(status_code=500, detail="Database connection failed")

    cursor = conn.cursor()

    try:
        for row in records:
            if not row.get("name"):
                continue

            image_name = row.get("image_name") or row.get("image_path") or row.get("image")
            image_path = None
            
            if image_name and os.path.exists(images_folder):
                local_img_path = os.path.join(images_folder, image_name)
                if os.path.exists(local_img_path):
                    try:
                        with open(local_img_path, "rb") as img_file:
                            img_content = img_file.read()
                        
                        pil_image = Image.open(io.BytesIO(img_content))
                        processed_image = crop_and_resize_to_square(pil_image, 500)
                        webp_bytes = compress_to_webp(processed_image)
                        
                        filename = f"{uuid.uuid4().hex}.webp"
                        storage_path = f"players/{filename}"
                        image_path = upload_image_to_supabase(webp_bytes, storage_path)
                    except Exception as e:
                        print(f"❌ Failed to process zip image {image_name}: {e}")

            cursor.execute("""
                INSERT INTO players (
                    name, nickname, age, gender, category, jersey, type,
                    mobile_no, email_id, base_price,
                    total_runs, highest_runs, wickets_taken,
                    times_out, teams_played, image_path
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (name) DO NOTHING
            """, (
                row.get("name"),
                row.get("nickname"),
                row.get("age"),
                row.get("gender"),
                row.get("category"),
                row.get("jersey"),
                row.get("type"),
                row.get("mobile_no") or row.get("mobile_No") or row.get("mobile"),
                row.get("email_id") or row.get("email_Id") or row.get("email"),
                row.get("base_price"),
                row.get("total_runs"),
                row.get("highest_runs"),
                row.get("wickets_taken"),
                row.get("times_out"),
                row.get("teams_played"),
                image_path
            ))

        conn.commit()

        return {"message": "ZIP upload successful 🚀"}

    except Exception as e:
        conn.rollback()
        print("❌ upload error:", e)
        raise HTTPException(500, str(e))

    finally:
        cursor.close()
        conn.close()
        if os.path.exists(temp_dir):
            try:
                shutil.rmtree(temp_dir)
            except Exception as rmtree_err:
                print(f"⚠️ Error cleaning up temp dir: {rmtree_err}")
